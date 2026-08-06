from openpyxl import load_workbook

archivo = r"C:\Users\OSCAR DAMIAN\OneDrive\Escritorio\AuditoriaFiscal\LIBRO MAYOR DIKEN DE MEXICO MAY26.xlsx"

try:
    wb = load_workbook(
        archivo,
        read_only=True,
        data_only=True
    )

    print("Hojas:")
    print(wb.sheetnames)

except Exception as e:
    print(type(e).__name__)
    print(e)